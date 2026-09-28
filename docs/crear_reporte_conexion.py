"""Genera el informe académico editable. Ejecutar con Python y python-docx."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'documentos' / 'Informe_conexion_y_cache_FiadOS.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.8)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    style = d.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
d.styles['Normal'].font.size = Pt(11)
d.styles['Normal'].paragraph_format.space_after = Pt(7)
d.styles['Normal'].paragraph_format.line_spacing = 1.08
d.styles['Title'].font.size = Pt(24)
d.styles['Heading 1'].font.size = Pt(16)
d.styles['Heading 2'].font.size = Pt(12)
for name in ('Heading 1', 'Heading 2'):
    d.styles[name].paragraph_format.space_before = Pt(10)
    d.styles[name].paragraph_format.space_after = Pt(5)

def p(text):
    return d.add_paragraph(text)
def h(text, level=1):
    return d.add_heading(text, level=level)
def page():
    d.add_page_break()
def tabla(headers, rows, widths):
    t = d.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in zip(t.columns, widths): c.width = Inches(w)
    for c, val in zip(t.rows[0].cells, headers): c.text = val
    for vals in rows:
        for c, val in zip(t.add_row().cells, vals): c.text = val
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        for j, c in enumerate(row.cells):
            c.width = Inches(widths[j]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '334155' if i == 0 else ('F1F5F9' if i % 2 == 0 else 'FFFFFF')); pr.append(shade)
            borders = OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e = OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                e = OxmlElement('w:'+edge); e.set(qn('w:w'),'100'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.line_spacing = 1
                for run in par.runs:
                    run.font.size = Pt(10)
                    if i == 0: run.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    p('')

d.add_paragraph('Conexión y almacenamiento temporal en FiadOS', 'Title')
p('Informe de implementación y pruebas • 28 de septiembre de 2026')
p('Alumno ____________________    Grupo __________    Materia ____________________')
h('Objetivo y resultado')
p('Se implementó una estrategia básica para reaccionar ante fallos de red o de acceso a los datos. La aplicación muestra el estado de conexión, comunica errores y recupera consultas guardadas cuando la API no responde. Las pruebas automáticas simularon la desconexión y comprobaron la recuperación de la información.')
p('El alcance es consulta sin conexión con la aplicación ya abierta y una sesión existente. Crear clientes, fiados o abonos requiere confirmación del servidor. No se implementó una cola de registros pendientes ni sincronización automática. Las capturas de evidencia se incorporarán en las últimas páginas.')
h('Cómo se detecta la conexión')
p('ConexionService escucha los eventos online y offline y consulta navigator.onLine. Esta señal informa sobre la red del navegador, pero no garantiza que PHP o MySQL funcionen. Por eso también se observa el resultado de las peticiones y se ofrece el botón Comprobar conexión, que consulta estado.php sin devolver datos del negocio.')
p('Si se desconecta el Wi-Fi pero XAMPP sigue activo en la misma computadora, localhost puede continuar respondiendo. La aplicación no bloquea esas peticiones solo porque el navegador indique falta de red.')
h('Estrategia de caché utilizada')
p('Se utilizó localStorage del navegador con datos JSON y estrategia de red primero. Cada consulta intenta obtener datos actuales; si tiene éxito, guarda una copia. Ante fallo de transporte, tiempo de espera agotado o error 500 del servidor, usa la copia vigente de esa misma consulta y muestra su fecha.')
tabla(['Regla', 'Implementación'], [
 ('Datos incluidos', 'Clientes, detalles, dashboard e historiales previamente consultados.'),
 ('Identificación', 'Negocio + usuario + URL completa, incluidos filtros e identificadores.'),
 ('Vigencia y espacio', '24 horas; hasta 30 consultas y aproximadamente 1 MB de texto.'),
 ('Limpieza', 'Al entrar o cerrar sesión y después de un cambio confirmado por la API.'),
 ('Datos excluidos', 'Contraseñas, nuevos registros pendientes, PDF y configuración administrativa.')
], [1.55, 5.35])

page(); h('Bitácora de problemas y soluciones')
p('Los siguientes problemas se identificaron en el código y durante la implementación. Se utilizó asistencia de IA para revisar el flujo, proponer cambios, escribir pruebas y elaborar este informe. Los resultados se comprobaron con ejecución de pruebas y compilación.')
h('Problema 1 Consultas dependientes del servidor', 2)
p('Antes: ClientesService y MovimientosService consultaban la API sin conservar respuestas para recuperarlas después de un fallo. Al perder acceso a los datos, solo quedaba un mensaje de error.')
p('Solución: CacheService guarda respuestas exitosas en localStorage; resilienciaInterceptor recupera la copia si ocurre un fallo temporal. Validación: una consulta exitosa seguida de un error de transporte devuelve el mismo cliente de prueba y activa el aviso de consulta local.')
h('Problema 2 Mensajes genéricos y falta de estado visible', 2)
p('Antes: varias pantallas mostraban Error al conectar con el servidor y Clientes indicaba revisar XAMPP, sin distinguir permisos, sesión expirada o falta de una copia local.')
p('Solución: se agregó un indicador global y mensajes específicos para errores de transporte, 401, 403, 500 y espera agotada. Las pantallas principales usan el mensaje del interceptor. Validación: las pruebas comprobaron mensajes, estado de red y respuestas de permisos; el navegador mostró Servidor disponible al comprobar la API local.')
h('Problema 3 Peticiones sin tiempo máximo explícito', 2)
p('Antes: las llamadas HTTP no fijaban un límite de espera de la aplicación. Un servidor que no respondía podía mantener la interfaz esperando demasiado tiempo.')
p('Solución: límite de 10 segundos para peticiones normales y 60 segundos para reportes PDF. Validación: con reloj simulado, una consulta que no respondió se canceló al superar 10 segundos y recuperó su copia. Las escrituras no se reintentan automáticamente.')
h('Problema 4 Riesgo de usar datos antiguos o de otra sesión', 2)
p('Durante el diseño de la caché se identificó que una clave compartida podía mezclar usuarios o conservar saldos anteriores a un pago. También una petición iniciada antes de guardar podía terminar después y restaurar una copia antigua.')
p('Solución: claves por negocio y usuario, vencimiento, limpieza al cerrar sesión y al confirmar cambios, y control de revisión para respuestas anteriores. Validación: pruebas de separación, limpieza, caducidad y respuesta tardía. Los errores 401 y 403 nunca se sustituyen con datos de caché.')
h('Problema 5 Detalles internos en errores de MySQL', 2)
p('Antes: config.php devolvía al navegador el texto técnico de la excepción de conexión. Solución: se registra el detalle en el servidor y se responde con un mensaje general. Se verificó con una API de prueba configurada para una base inexistente, sin modificar la base real.')

page(); h('Pruebas y resultados')
p('Se ejecutaron 20 pruebas de Angular y Vitest: las 5 existentes y 15 nuevas para conexión y caché. Todas pasaron. Las nuevas usan HttpTestingController, que simula fallos del transporte y respuestas HTTP sin crear clientes ni movimientos reales. También se compiló la aplicación para producción y se validó la sintaxis de PHP.')
tabla(['Caso', 'Resultado comprobado'], [
 ('P01 Consulta y desconexión', 'Se devuelve la copia guardada y se marca servidor inaccesible.'),
 ('P02 Persistencia', 'Una nueva instancia de CacheService recupera el JSON de localStorage.'),
 ('P03 Wi-Fi y localhost', 'El evento offline cambia el indicador; una API que responde sigue siendo usable.'),
 ('P04 Sin copia', 'Se informa que no existe copia vigente; no se inventan resultados.'),
 ('P05 Caducidad', 'Se rechaza una copia con más de 24 horas.'),
 ('P06 Separación', 'Otro negocio u otro usuario no recupera esa copia.'),
 ('P07 Sesión reemplazada', 'Una respuesta iniciada con otro token no se guarda.'),
 ('P08 y P09 Permisos', '401 elimina sesión y caché; 403 se informa sin usar caché.'),
 ('P10 Recuperación', 'Un error 500 usa caché; una respuesta nueva la reemplaza y retira su aviso.'),
 ('P11 Espera agotada', 'La consulta se cancela a los 10 segundos y recupera la copia.'),
 ('P12 Escritura fallida', 'El abono no se presenta como confirmado ni se reenvía automáticamente.'),
 ('P13 y P14 Limpieza', 'Cambios confirmados y cierre de sesión eliminan copias; una lectura anterior no las restaura.'),
 ('P15 Almacenamiento', 'JSON corrupto o cuota agotada no impiden consultar la API; se muestra advertencia.')
], [2.05, 4.85])
h('Archivos principales', 2)
p('src/app/services/cache.service.ts: guarda, obtiene y limpia copias. conexion.service.ts: mantiene señales de red y API. resiliencia.interceptor.ts: tiempo de espera, errores y recuperación. conexion-banner.component.ts: mensajes visibles. resiliencia.interceptor.spec.ts: pruebas reproducibles. api/dermexcel-api/estado.php: comprobación mínima del servicio.')
h('Límites de la solución', 2)
p('La caché pertenece al navegador y al origen donde se abrió la aplicación. No es un respaldo de MySQL y no está cifrada. Solo recupera URLs ya consultadas; otro filtro puede no tener copia. No guarda el programa para iniciarlo desde cero sin red. Una escritura cuyo resultado no llegó puede haberse procesado: se debe revisar el historial antes de repetirla.')

page(); h('Guía para agregar evidencias de consulta')
p('Realizar las capturas en Chrome o Edge con la aplicación abierta. Abrir primero Inicio, Clientes y el detalle de un cliente para llenar la caché. Mantener la sesión iniciada. Usar datos de prueba y ocultar información personal en las capturas.')
h('Captura 1 Consulta con conexión', 2)
p('Abrir Clientes y esperar la lista. Pulsar Comprobar conexión. La captura debe mostrar Servidor disponible y los clientes. En F12, Application o Aplicación, Local Storage, puede comprobarse la clave fiados.consultas.v1.')
p('Insertar aquí la captura con conexión.').paragraph_format.space_after = Pt(105)
p('Fecha y hora ____________________    Observaciones __________________________')
h('Captura 2 Consulta sin acceso a la API', 2)
p('En F12, abrir Network o Red y seleccionar Offline o Sin conexión. Sin recargar la pestaña, ir a Inicio y regresar a Clientes. La petición fallará y se debe mostrar Consulta local, su fecha y los clientes guardados. No basta con desconectar Wi-Fi si la API está en localhost: el servicio local podría seguir funcionando.')
p('Insertar aquí la captura del aviso de consulta local.').paragraph_format.space_after = Pt(105)
p('Fecha y hora ____________________    Observaciones __________________________')
p('Si el navegador permite localhost aun en modo Offline, usar el bloqueo de solicitudes de DevTools para el patrón *dermexcel-api/*. Desactivar ese bloqueo al terminar. No recargar toda la aplicación durante la prueba.')

page(); h('Guía para agregar evidencias de errores')
h('Captura 3 Registro no confirmado', 2)
p('Con las solicitudes de la API bloqueadas y un cliente de prueba previamente cargado, abrir Registrar pago, introducir un monto válido y pulsar guardar. La captura debe mostrar que no se pudo confirmar el registro y no debe aparecer Pago registrado correctamente. No realizar esta prueba con dinero real ni enviar el registro después de reconectar.')
p('Insertar aquí la captura del error de registro.').paragraph_format.space_after = Pt(110)
p('Fecha y hora ____________________    Observaciones __________________________')
h('Captura 4 Recuperación de la conexión', 2)
p('Cambiar Offline a No throttling o Sin limitación y quitar cualquier bloqueo de solicitudes. Pulsar Comprobar conexión. Luego cambiar de pestaña y volver a Clientes para volver a consultar la API. La información de esa consulta se actualiza y su aviso de copia se retira. Si quedan otras consultas antiguas, volver a abrir esas pantallas para actualizarlas.')
p('Insertar aquí la captura de recuperación.').paragraph_format.space_after = Pt(110)
p('Fecha y hora ____________________    Observaciones __________________________')
h('Conclusión', 2)
p('La aplicación reacciona a fallos de red y de datos con avisos y consultas locales controladas. MySQL continúa siendo la fuente oficial. La solución permite consultar información previamente cargada durante un corte, pero requiere el servidor para confirmar nuevos registros. Las capturas complementan las pruebas automáticas con evidencia visual de la ejecución del alumno.')

d.core_properties.title = 'Conexión y almacenamiento temporal en FiadOS'
d.core_properties.subject = 'Implementación, bitácora y pruebas de consulta sin conexión'
d.core_properties.author = ''
for root in (d._element, d.styles.element):
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
d.save(OUT)
print(OUT)
